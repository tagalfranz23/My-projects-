"""Editable Microsoft Word document generation for barangay records.

The functions in this module return seeked ``BytesIO`` streams suitable for
Flask's ``send_file``.  All content is created as native DOCX paragraphs and
tables so authorized users can edit the downloaded files in Microsoft Word or
another compatible editor.
"""

from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCX_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_NAVY = "082F5B"
_GOLD = "D4A62A"
_PALE_BLUE = "EAF1F8"
_LIGHT_GRAY = "F3F4F6"


def _text(value, fallback="-"):
    if value is None or value == "":
        return fallback
    if isinstance(value, datetime):
        return value.strftime("%B %d, %Y at %I:%M %p")
    if isinstance(value, date):
        return value.strftime("%B %d, %Y")
    if isinstance(value, time):
        return value.strftime("%I:%M %p")
    return str(value)


def _money(value):
    if value is None or value == "":
        return None
    return f"PHP {Decimal(str(value)):,.2f}"


def _set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_cell_text(cell, value, *, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(_text(value))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _prepare_document(title, subject, barangay_name):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    properties = document.core_properties
    properties.title = title
    properties.subject = subject
    properties.author = barangay_name
    properties.keywords = "Barangay Minante 1, editable Word document"
    return document


def _add_barangay_header(document, barangay_name, location, document_title):
    for line, size, bold, color in (
        ("REPUBLIC OF THE PHILIPPINES", 9, False, None),
        (barangay_name.upper(), 15, True, _NAVY),
        (location, 9, False, None),
        ("BARANGAY MINANTE 1 INTEGRATED MANAGEMENT SYSTEM", 9, True, _NAVY),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(8)
    run = heading.add_run(document_title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(_NAVY)


def _add_reference_box(document, reference):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_cell_shading(table.cell(0, 0), _PALE_BLUE)
    _set_cell_shading(table.cell(0, 1), _PALE_BLUE)
    _set_cell_text(table.cell(0, 0), "REFERENCE NUMBER", bold=True, color=_NAVY)
    _set_cell_text(table.cell(0, 1), reference, bold=True, color=_NAVY, size=12)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_details_table(document, rows):
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        if value is None or value == "":
            continue
        cells = table.add_row().cells
        _set_cell_shading(cells[0], _LIGHT_GRAY)
        _set_cell_text(cells[0], label, bold=True, color=_NAVY)
        _set_cell_text(cells[1], value)
    return table


def _add_labeled_paragraph(document, heading, body):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(heading.upper())
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(_NAVY)
    body_paragraph = document.add_paragraph(_text(body, ""))
    body_paragraph.paragraph_format.space_after = Pt(3)


def _to_stream(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def _event_venue_name(record):
    venue = getattr(record, "venue", None)
    reservation = getattr(record, "reservation", None)
    if venue is None and reservation is not None:
        venue = getattr(reservation, "venue", None)
    return getattr(venue, "name", None) or getattr(record, "proposed_location", None)


def _event_schedule(record):
    event_date = getattr(record, "proposed_date", None)
    start = getattr(record, "proposed_start_time", None)
    end = getattr(record, "proposed_end_time", None)
    if event_date and start and end:
        return f"{_text(event_date)}, {_text(start)} to {_text(end)}"
    return _text(event_date)


def generate_submission_acknowledgment(
    record,
    submission_type,
    processing_time,
    barangay_name="Barangay Minante 1",
    location="Cauayan City, Isabela",
):
    """Return an editable submission acknowledgment as a DOCX stream."""
    kind = submission_type.lower().strip()
    if kind not in {"permit", "event", "blotter"}:
        raise ValueError("Unsupported acknowledgment type")

    if kind == "permit":
        owner = record.applicant
        submitted = record.application_date or record.created_at
        reference = record.reference_no
        label = "Permit Application"
        fee = _money(getattr(record, "fee_at_submission", None))
        details = [
            ("Permit Type", record.permit_type.name),
            ("Purpose", record.purpose),
            ("Fee at Submission", fee),
            ("Requirements / What to Bring", record.permit_type.requirements),
        ]
    elif kind == "event":
        owner = record.requester
        submitted = record.created_at
        reference = record.reference_no
        label = "Event Approval Request"
        category = getattr(record, "category", None)
        event_type = getattr(category, "name", None) or getattr(record, "event_type", None)
        fee = _money(getattr(record, "fee_at_submission", None))
        details = [
            ("Event Name", record.event_name),
            ("Event Category", event_type),
            ("Proposed Schedule", _event_schedule(record)),
            ("Venue", _event_venue_name(record)),
            ("Fee at Submission", fee),
        ]
    else:
        owner = record.complainant
        submitted = record.filed_date or record.created_at
        reference = record.case_no
        label = "Blotter Report / Complaint"
        details = [
            ("Incident Type", record.incident_type),
            ("Incident Date", record.incident_date),
            ("Incident Location", record.incident_location),
            ("Respondent / Person Complained Against", record.respondent_name),
        ]

    document = _prepare_document(
        f"Submission Acknowledgment {reference}",
        "Proof that Barangay Minante 1 received a submission",
        barangay_name,
    )
    _add_barangay_header(document, barangay_name, location, "SUBMISSION ACKNOWLEDGMENT")
    _add_reference_box(document, reference)
    _add_details_table(
        document,
        [
            ("Submission Type", label),
            ("Resident / Applicant", owner.full_name),
            ("Date and Time Submitted", submitted),
            ("Current Status", getattr(record, "status", None)),
            *details,
        ],
    )

    _add_labeled_paragraph(
        document,
        "Received Confirmation",
        (
            f"Barangay Minante 1 successfully received this submission on "
            f"{_text(submitted)}. Keep this acknowledgment and reference number "
            "for tracking and follow-up."
        ),
    )
    if processing_time:
        _add_labeled_paragraph(
            document,
            "Estimated Processing Time",
            (
                f"{processing_time}. This estimate is subject to completeness, "
                "verification, and Barangay review."
            ),
        )
    _add_labeled_paragraph(
        document,
        "Next Steps",
        (
            "Barangay personnel will review the submission. Additional information "
            "may be requested when necessary, and status updates will appear through "
            "the system's authorized notification channels."
        ),
    )

    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_before = Pt(8)
    run = notice.add_run(
        "This editable document confirms receipt of the submission only. It is not "
        "proof of payment, approval, authorization, resolution, permit issuance, or "
        "a final Barangay decision."
    )
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(139, 47, 47)

    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"Generated from authorized Barangay records | {reference}").font.size = Pt(8)
    return _to_stream(document)


def generate_walkin_claim_stub(application, queue_entry, barangay_name="Barangay Minante 1"):
    """Generate a filing/queue stub; it is never a payment receipt or Permit."""
    document = _prepare_document(
        f"Walk-in Claim Stub {application.reference_no}",
        "Walk-in filing and queue reference",
        barangay_name,
    )
    _add_barangay_header(document, barangay_name, "Cauayan City, Isabela", "WALK-IN CLAIM STUB")
    _add_reference_box(document, application.reference_no)
    _add_details_table(document, [
        ("Resident", application.applicant.full_name),
        ("Service", application.permit_type.name),
        ("Queue Number", queue_entry.queue_number),
        ("Fee at Submission", _money(application.fee_at_submission)),
        ("Filing Status", application.status),
        ("Important", "This is a claim stub and proof of filing. It is not a payment receipt or approved Permit."),
    ])
    return _to_stream(document)


def generate_approved_permit_document(
    application,
    applicant,
    permit_type,
    barangay_name="Barangay Minante 1",
    city="Cauayan City, Isabela",
    signatory_name=None,
    signatory_title=None,
    signature_file=None,
):
    """Return an approved permit as an editable DOCX stream."""
    reference = application.reference_no
    document = _prepare_document(
        f"Approved Permit {reference}",
        "Approved Barangay permit document",
        barangay_name,
    )
    _add_barangay_header(document, barangay_name, city, permit_type.name.upper())
    _add_reference_box(document, reference)

    statement = document.add_paragraph()
    statement.paragraph_format.space_before = Pt(12)
    statement.paragraph_format.space_after = Pt(12)
    statement.add_run("This is to certify that ")
    name_run = statement.add_run(applicant.full_name)
    name_run.bold = True
    statement.add_run(
        f", a resident of {_text(getattr(applicant, 'address', None), 'the address recorded by the Barangay')}, "
        f"is granted this {permit_type.name} for the purpose stated below."
    )

    fee = _money(getattr(application, "fee_at_submission", None))
    validity_days = getattr(permit_type, "validity_days", None)
    validity = f"{validity_days} days from issuance" if validity_days else None
    _add_details_table(
        document,
        [
            ("Reference Number", reference),
            ("Applicant", applicant.full_name),
            ("Purpose", application.purpose),
            ("Date Issued", application.decision_date),
            ("Validity", validity),
            ("Recorded Fee", fee),
            ("Status", application.status),
            ("Remarks", getattr(application, "remarks", None)),
        ],
    )

    document.add_paragraph()
    signatures = document.add_table(rows=3 if signature_file else 2, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    signatures.cell(0, 0).text = "____________________________"
    signatures.cell(0, 1).text = "____________________________"
    signatures.cell(1, 0).text = "Applicant's Signature"
    signatures.cell(1, 1).text = _text(signatory_title, "Authorized Barangay Official")
    if signature_file:
        try:
            signature_paragraph = signatures.cell(2, 1).paragraphs[0]
            signature_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            signature_paragraph.add_run().add_picture(signature_file, width=Inches(1.35))
            signatures.cell(2, 0).text = ""
        except (OSError, ValueError):
            # The signed record remains traceable even if an image is later unavailable.
            signatures.cell(2, 1).text = "Electronic signature on record"
    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    name.add_run(_text(signatory_name, "Authorized official")).bold = True
    name.add_run("\nElectronic signature applied by the authorized Barangay official.")
    for row in signatures.rows:
        for cell in row.cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
    return _to_stream(document)


def generate_system_report(
    title,
    headers,
    rows,
    subtitle=None,
    barangay_name="Barangay Minante 1",
):
    """Return an editable tabular system report as a DOCX stream."""
    document = _prepare_document(title, "Barangay system report", barangay_name)
    _add_barangay_header(
        document,
        barangay_name,
        "Cauayan City, Isabela",
        title.upper(),
    )
    if subtitle:
        paragraph = document.add_paragraph(_text(subtitle))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(10)

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_row = table.rows[0]
    _set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        _set_cell_shading(cell, _NAVY)
        _set_cell_text(cell, header, bold=True, color="FFFFFF")

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if index >= len(cells):
                break
            _set_cell_text(cells[index], value, size=8)

    if not rows:
        empty = table.add_row().cells
        _set_cell_text(empty[0], "No records matched the selected report criteria.")
        if len(empty) > 1:
            empty[0].merge(empty[-1])

    footer = document.add_paragraph()
    footer.paragraph_format.space_before = Pt(8)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Editable Word report generated from authorized system records.")
    run.italic = True
    run.font.size = Pt(8)
    return _to_stream(document)
